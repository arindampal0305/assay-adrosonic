"""File ingestion (FR-ING-01 to FR-ING-05).

Reads .xlsx through openpyxl so merged-cell ranges survive, and .csv through the
stdlib reader so ragged title rows do not break tokenisation. Merged ranges are
forward-filled before any analysis, which is what turns a banner spanning three
columns into three composite headers.
"""

from __future__ import annotations

import csv
import hashlib
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Optional

import openpyxl
import pandas as pd

MAX_DATA_ROWS = 5000
SUPPORTED_SUFFIXES = {".xlsx", ".csv"}


class IngestError(Exception):
    """Unsupported, empty, corrupt or oversized input. Carries a reader-friendly message."""


@dataclass
class MergedSpan:
    row_start: int
    col_start: int
    row_end: int
    col_end: int

    @property
    def width(self) -> int:
        return self.col_end - self.col_start + 1

    def covers_row(self, row: int) -> bool:
        return self.row_start <= row <= self.row_end


@dataclass
class LoadedSheet:
    name: str
    grid: list[list[Any]]
    merged_spans: list[MergedSpan] = field(default_factory=list)

    @property
    def n_rows(self) -> int:
        return len(self.grid)

    @property
    def n_cols(self) -> int:
        return len(self.grid[0]) if self.grid else 0

    def horizontal_spans_on_row(self, row: int) -> list[MergedSpan]:
        return [s for s in self.merged_spans if s.width > 1 and s.covers_row(row)]

    def as_dataframe(self, header_row: int, data_start_row: int, headers: list[str]) -> pd.DataFrame:
        body = self.grid[data_start_row:]
        return pd.DataFrame(body, columns=headers[: self.n_cols], dtype=object)


@dataclass
class LoadedWorkbook:
    file_name: str
    sha256: str
    sheets: list[LoadedSheet]


def _normalise_cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        stripped = value.strip()
        return stripped or None
    if isinstance(value, (datetime, date)):
        return value
    return value


def _trim_grid(rows: list[list[Any]]) -> list[list[Any]]:
    """Drop trailing empty rows and columns; openpyxl often overshoots max_row."""
    while rows and all(c is None for c in rows[-1]):
        rows.pop()
    if not rows:
        return []
    width = max(len(r) for r in rows)
    last_used = -1
    for row in rows:
        for idx in range(len(row) - 1, -1, -1):
            if row[idx] is not None:
                last_used = max(last_used, idx)
                break
    width = last_used + 1 if last_used >= 0 else width
    return [list(r[:width]) + [None] * (width - len(r[:width])) for r in rows]


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _apply_merges(rows: list[list[Any]], spans: list[MergedSpan]) -> None:
    for span in spans:
        if span.row_start >= len(rows):
            continue
        anchor_row = rows[span.row_start]
        if span.col_start >= len(anchor_row):
            continue
        value = anchor_row[span.col_start]
        if value is None:
            continue
        for r in range(span.row_start, min(span.row_end, len(rows) - 1) + 1):
            for c in range(span.col_start, min(span.col_end, len(rows[r]) - 1) + 1):
                rows[r][c] = value


def _load_xlsx(path: Path) -> list[LoadedSheet]:
    try:
        workbook = openpyxl.load_workbook(path, data_only=True, read_only=False)
    except zipfile.BadZipFile as exc:
        raise IngestError(
            "This .xlsx file could not be opened. It is either corrupt or "
            "password-protected. Please save an unprotected copy and upload again."
        ) from exc
    except Exception as exc:
        raise IngestError(f"This .xlsx file could not be read: {exc}") from exc

    sheets: list[LoadedSheet] = []
    try:
        for worksheet in workbook.worksheets:
            raw = [
                [_normalise_cell(v) for v in row]
                for row in worksheet.iter_rows(values_only=True)
            ]
            grid = _trim_grid(raw)
            spans = [
                MergedSpan(
                    row_start=rng.min_row - 1,
                    col_start=rng.min_col - 1,
                    row_end=rng.max_row - 1,
                    col_end=rng.max_col - 1,
                )
                for rng in worksheet.merged_cells.ranges
            ]
            _apply_merges(grid, spans)
            sheets.append(LoadedSheet(name=worksheet.title, grid=grid, merged_spans=spans))
    finally:
        workbook.close()
    return sheets


def _load_csv(path: Path, sheet_name: str) -> list[LoadedSheet]:
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        sample = fh.read(64 * 1024)
        fh.seek(0)
        try:
            dialect: Any = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        raw = [[_normalise_cell(v) for v in row] for row in csv.reader(fh, dialect)]
    return [LoadedSheet(name=sheet_name, grid=_trim_grid(raw))]


def load_workbook_file(path: Path | str, file_name: Optional[str] = None) -> LoadedWorkbook:
    path = Path(path)
    display_name = file_name or path.name
    suffix = Path(display_name).suffix.lower() or path.suffix.lower()

    if suffix not in SUPPORTED_SUFFIXES:
        raise IngestError(
            f"'{display_name}' has an unsupported type ({suffix or 'no extension'}). "
            "ASSAY accepts .xlsx and .csv only."
        )
    if not path.exists():
        raise IngestError(f"'{display_name}' could not be found on disk.")
    if path.stat().st_size == 0:
        raise IngestError(f"'{display_name}' is empty (0 bytes).")

    # A CSV has no sheet name of its own, so it takes the uploaded file's name
    # rather than the temporary path it was stored under.
    sheets = (
        _load_xlsx(path) if suffix == ".xlsx" else _load_csv(path, Path(display_name).stem)
    )
    sheets = [s for s in sheets if s.n_rows > 0]
    if not sheets:
        raise IngestError(f"'{display_name}' contains no readable rows on any sheet.")

    largest = max(s.n_rows for s in sheets)
    if largest > MAX_DATA_ROWS + 50:
        raise IngestError(
            f"'{display_name}' has {largest} rows on its largest sheet, above the "
            f"{MAX_DATA_ROWS}-row limit. Please split the file and upload the parts."
        )

    return LoadedWorkbook(file_name=display_name, sha256=_sha256(path), sheets=sheets)
