"""Generate practice SOVs in samples/.

These stand in for the three real test SOVs until they are released. They
reproduce the mess the SRS and planning document describe: shifted headers under
title blocks, merged banner headers, abbreviations, multi-sheet workbooks with
notes tabs, missing target columns, placeholders and currency strings.

Drop the real files into samples/ and run_samples.py picks them up unchanged.
"""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

SAMPLES = Path(__file__).resolve().parents[1] / "samples"

SAMPLE1_HEADERS = [
    "Loc #", "Street Address", "City", "St", "Zip", "Bldg Repl Cost", "Contents",
    "BI/EE", "Occ", "Const", "Stories", "# Bldgs", "Yr Blt", "Sprk",
]
SAMPLE1_ROWS = [
    ["LOC-001", "1420 Industrial Pkwy", "Austin", "TX", 78701, 1250000, 480000, 95000, "Warehouse", "Masonry", 1, 1, 1998, "Y"],
    ["LOC-002", "88 Harbor St", "Boston", "MA", 2110, 3400000, 1200000, 260000, "Office", "Fire Resistive", 8, 1, 1972, "Y13"],
    ["LOC-003", "9900 Rio Grande Blvd", "Albuquerque", "NM", 87114, 640000, 155000, 0, "Retail", "Frame", 1, 2, 2005, "N"],
    ["LOC-004", "17 Mill Race Rd", "Greenville", "SC", 29601, "$2,100,000.00", "$775,000.00", 180000, "Manufacturing", "Joisted Masonry", 2, 3, 1965, "Y"],
    ["LOC-005", "455 Lakeview Dr", "Madison", "WI", 53703, 890000, 210000, 45000, "Office", "Masonry", 3, 1, 0, "N"],
    ["LOC-006", "2 Cannery Row", "Monterey", "CA", 93940, 1575000, 640000, 130000, "Restaurant", "Frame", 1, 1, 1988, "Y(13R)"],
    ["LOC-007", "6100 Peachtree Ind", "Atlanta", "GA", 30341, -45000, 98000, 12000, "Storage", "Metal", 1, 1, 2012, "N"],
    ["LOC-008", "311 N Clark", "Chicago", "IL", 60654, 7800000, 2450000, 610000, "Office", "Fire Resistive", 24, 1, 1955, "Y"],
    ["LOC-009", "742 Evergreen Ter", "Springfield", "OR", 97477, 320000, 86000, 0, "Dwelling", "Frame", 2, 1, 1931, "N"],
    ["LOC-010", "1 Riverside Plz", "Columbus", "OH", 43215, 4100000, 1340000, 320000, "Office", "Masonry", 12, 1, 1979, "Y"],
    ["LOC-011", "500 W Monroe", "Phoenix", "AZ", 85003, 1980000, 505000, 110000, "Mixed Use", "Masonry", 4, 2, 2019, "Y13"],
    ["LOC-012", "27 Dock Rd", "Portland", "ME", 4101, 730000, 190000, 38000, "Marina", "Frame", 1, 4, 2031, "Partial"],
]

SAMPLE2_HEADERS = [
    "Location Number", "Property Address", "Town", "State Name", "Postal Code",
    "County", "Replacement Cost", "Business Personal Property", "Loss of Income",
    "Other Property", "Use", "Wall Type", "No of Floors", "Qty Buildings",
    "Year of Construction", "Auto Sprinkler",
]
SAMPLE2_ROWS = [
    [1001, "400 Commerce Way", "Reno", "Nevada", 89501, "Washoe", "$4,250,000.00", "$1,100,000.00", "$320,000.00", 15000, "Distribution", "Tilt-Up Concrete", 1, 2, 2001, "Yes"],
    [1002, "22 Beacon Hill Rd", "Hartford", "Connecticut", 6103, "Hartford", "$2,875,000.00", "$940,000.00", "$210,000.00", 0, "Office", "Masonry", 6, 1, 1968, "Yes"],
    [1003, "7781 Gulf Breeze Hwy", "Pensacola", "Florida", 32507, "Escambia", "$1,340,000.00", "$385,000.00", "$75,000.00", 8000, "Retail", "Frame", 1, 1, 1994, "No"],
    [1004, "15 Granite Quarry Ln", "Barre", "Vermont", 5641, "Washington", "$615,000.00", "$142,000.00", "$0.00", 0, "Workshop", "Metal", 1, 3, 1900, "No"],
    [1005, "3030 Airport Fwy", "Fort Worth", "Texas", 76111, "Tarrant", "$9,100,000.00", "$3,200,000.00", "$880,000.00", 42000, "Manufacturing", "Fire Resistive", 3, 5, 1983, "13R"],
    [1006, "601 Union St", "Seattle", "Washington", 98101, "King", "$12,400,000.00", "$4,050,000.00", "$1,150,000.00", 65000, "Office", "Fire Resistive", 31, 1, 1990, "Yes"],
    [1007, "88 Shoreline Dr", "Duluth", "Minnesota", 55802, "St. Louis", "$1,020,000.00", "$265,000.00", "$48,000.00", 0, "Warehouse", "Joisted Masonry", 2, 1, 1900, "No"],
    [1008, "4 Canyon Rim Ct", "Boise", "Idaho", 83702, "Ada", "$845,000.00", "$198,000.00", "$36,000.00", 5000, "Office", "Frame", 2, 1, 2015, "Yes"],
]

SAMPLE3_TOP = [
    "Loc Ref", "Street Address", "City", "St", "Zip", "Values", None, None,
    "Occ", "Const Type", "Stories", "# Bldgs", "Yr Blt", "Sprk",
]
SAMPLE3_SUB = [
    None, None, None, None, None, "Building", "Contents", "BI/EE",
    None, None, None, None, None, None,
]
SAMPLE3_ROWS = [
    ["L-01", "250 Foundry St", "Pittsburgh", "PA", 15222, 5600000, 1750000, 420000, "Manufacturing", "Joisted Masonry", 2, 4, 1948, "Y"],
    ["L-02", "1 Allegheny Ctr", "Pittsburgh", "PA", 15212, 8900000, 2300000, 640000, "Office", "Fire Resistive", 18, 1, 1976, "Y"],
    ["L-03", "4455 Steubenville Pike", "Robinson", "PA", 15205, 1450000, 410000, 88000, "Retail", "Masonry", 1, 1, 2003, "N"],
    ["L-04", "77 Mon Wharf", "Pittsburgh", "PA", 15219, 980000, 0, 0, "Parking", "Concrete", 1, 1, 1900, "N"],
    ["L-05", "3 Cranberry Woods Dr", "Cranberry Twp", "PA", 16066, 3250000, 1120000, 295000, "Office", "Tilt-Up Concrete", 4, 2, 2008, "Y13"],
    ["L-06", "920 Ohio River Blvd", "Sewickley", "PA", 15143, 672000, 158000, 31000, "Warehouse", "Metal", 1, 1, 1991, "N"],
    ["L-07", "1600 Smallman St", "Pittsburgh", "PA", 15222, 2140000, 890000, 175000, "Mixed Use", "Masonry", 3, 1, 1929, "Y(13R)"],
]

SAMPLE3_PRIOR_HEADERS = ["Loc Ref", "Street Address", "City", "St", "Building Value", "Contents", "BI", "Occupancy"]
SAMPLE3_PRIOR_ROWS = [
    ["L-01", "250 Foundry St", "Pittsburgh", "PA", 5400000, 1690000, 405000, "Manufacturing"],
    ["L-02", "1 Allegheny Ctr", "Pittsburgh", "PA", 8600000, 2240000, 620000, "Office"],
    ["L-03", "4455 Steubenville Pike", "Robinson", "PA", 1400000, 398000, 85000, "Retail"],
    ["L-04", "77 Mon Wharf", "Pittsburgh", "PA", 950000, 0, 0, "Parking"],
    ["L-05", "3 Cranberry Woods Dr", "Cranberry Twp", "PA", 3150000, 1080000, 286000, "Office"],
]


def _write_rows(ws, rows, start_row=1):
    for r, row in enumerate(rows, start=start_row):
        for c, value in enumerate(row, start=1):
            if value is not None:
                ws.cell(row=r, column=c, value=value)


def build_sample1() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    _write_rows(ws, [SAMPLE1_HEADERS] + SAMPLE1_ROWS)
    path = SAMPLES / "sample1_basic.xlsx"
    wb.save(path)
    return path


def build_sample1_csv() -> Path:
    path = SAMPLES / "sample1_basic.csv"
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(SAMPLE1_HEADERS)
        writer.writerows(SAMPLE1_ROWS)
    return path


def build_sample2() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "SOV"
    ws.cell(row=1, column=1, value="STATEMENT OF VALUES")
    ws.cell(row=2, column=1, value="Broker: Hacksmiths Risk Services")
    ws.cell(row=2, column=4, value="Valuation Date: 01/01/2026")
    _write_rows(ws, [SAMPLE2_HEADERS] + SAMPLE2_ROWS, start_row=4)
    path = SAMPLES / "sample2_title_block.xlsx"
    wb.save(path)
    return path


def build_sample3() -> Path:
    wb = Workbook()

    instructions = wb.active
    instructions.title = "Instructions"
    for r, text in enumerate(
        [
            "Please complete all columns on the Location Schedule tab.",
            "Values should be reported at 100% replacement cost.",
            "Return the completed workbook to your broker 30 days before renewal.",
            "Questions: risk.services@example.com",
        ],
        start=1,
    ):
        instructions.cell(row=r, column=1, value=text)

    schedule = wb.create_sheet("Location Schedule")
    schedule.cell(row=1, column=1, value="ACME Manufacturing — Location Schedule FY2026")
    schedule.merge_cells(
        start_row=1, start_column=1, end_row=1, end_column=len(SAMPLE3_TOP)
    )
    _write_rows(schedule, [SAMPLE3_TOP, SAMPLE3_SUB] + SAMPLE3_ROWS, start_row=3)
    schedule.merge_cells(start_row=3, start_column=6, end_row=3, end_column=8)

    prior = wb.create_sheet("Prior Year Values")
    _write_rows(prior, [SAMPLE3_PRIOR_HEADERS] + SAMPLE3_PRIOR_ROWS)

    notes = wb.create_sheet("Notes")
    for r, (label, value) in enumerate(
        [
            ("Prepared by", "J. Okafor"),
            ("Reviewed by", "M. Lindqvist"),
            ("Date", "2026-09-14"),
            ("Contact", "+1 412 555 0133"),
        ],
        start=1,
    ):
        notes.cell(row=r, column=1, value=label)
        notes.cell(row=r, column=2, value=value)

    path = SAMPLES / "sample3_multisheet_merged.xlsx"
    wb.save(path)
    return path


def main() -> None:
    SAMPLES.mkdir(parents=True, exist_ok=True)
    for builder in (build_sample1, build_sample1_csv, build_sample2, build_sample3):
        print("wrote", builder().name)


if __name__ == "__main__":
    main()
